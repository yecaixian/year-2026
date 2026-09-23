"""Step 3: 生成 XLSX 工作簿 (多工作表 + 格式)。"""
import csv
import json
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ANN = Path(r"C:\Users\lenovo\WorkBuddy\2026-09-21-20-29-42\data\annotations")
OUT = ANN / "Dseed_annotation_table.xlsx"

DIMS_CN = ["色彩协调性", "构图新颖性", "主题清晰度", "感知技术复杂度", "整体美学偏好"]
CLASSES_CN = ["GAN 合成", "噪声场", "像素拼贴", "光流", "故障", "粒子",
              "程序化纹理", "数据映射", "体素", "着色器", "分形", "算法笔触"]

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
SUB_FILL = PatternFill("solid", fgColor="D9E2F3")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws, row=1, ncol=None):
    ncol = ncol or ws.max_column
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def autosize(ws, minto=8, maxto=26):
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        w = minto
        for row in range(1, min(ws.max_row, 200) + 1):
            v = ws.cell(row=row, column=col).value
            if v is None:
                continue
            ln = sum(2 if ord(ch) > 127 else 1 for ch in str(v))
            w = max(w, min(maxto, ln + 3))
        ws.column_dimensions[letter].width = w


def read(path):
    with open(path, encoding="utf-8-sig") as f:
        return list(csv.reader(f))


def main():
    meta = json.loads((ANN / "Dseed_meta.json").read_text(encoding="utf-8"))
    main_tbl = read(ANN / "Dseed_annotation_table.csv")
    full_tbl = read(ANN / "Dseed_annotation_full.csv")
    raw_tbl = read(ANN / "Dseed_raw_annotators.csv")

    wb = Workbook()

    # ---------- Sheet 1: 标注总表 ----------
    ws = wb.active
    ws.title = "标注总表"
    for r in main_tbl:
        ws.append(r)
    style_header(ws)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=6):
        for cell in row:
            cell.number_format = "0.00"
            cell.alignment = Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = BORDER
    autosize(ws)

    # ---------- Sheet 2: 扩展字段 ----------
    ws2 = wb.create_sheet("扩展字段")
    for r in full_tbl:
        ws2.append(r)
    style_header(ws2)
    ws2.freeze_panes = "B2"
    ws2.auto_filter.ref = f"A1:{get_column_letter(ws2.max_column)}{ws2.max_row}"
    for row in ws2.iter_rows(min_row=2, min_col=2, max_col=6):
        for cell in row:
            cell.number_format = "0.00"
            cell.alignment = Alignment(horizontal="center")
    autosize(ws2)

    # ---------- Sheet 3: 标注者原始评分 ----------
    ws3 = wb.create_sheet("标注者原始评分")
    for r in raw_tbl:
        ws3.append(r)
    style_header(ws3)
    ws3.freeze_panes = "A2"
    autosize(ws3)

    # ---------- Sheet 4: 摘要与分布 ----------
    ws4 = wb.create_sheet("摘要与分布")
    ws4.append(["指标", "取值"])
    style_header(ws4)
    rows4 = [
        ("图像总数 N", meta["n_images"]),
        ("风格类别数", meta["n_classes"]),
        ("标注者人数", meta["n_annotators"]),
        ("评价维度数", 5),
        ("评分量表", "1–5 级李克特量表"),
        ("每图最终评分", "3 位标注者评分的算术平均值"),
        ("随机种子 (可复现)", meta["seed"]),
        ("标注者噪声 σ (模拟参数)", meta["annotator_noise_sigma"]),
        ("风格标签数", len(meta["class_distribution"])),
        ("_来源构成", json.dumps(meta["source_dataset_mix"], ensure_ascii=False)),
    ]
    for r in rows4:
        ws4.append(list(r))
    ws4.append([])
    ws4.append(["— 标注者间一致性 Krippendorff's α (区间型，逐维度估计) —"])
    ws4.append(["维度", "α", "95% CI 下界", "95% CI 上界"])
    for i, d in enumerate(["colour_harmony", "composition_novelty", "theme_clarity",
                           "technical_complexity", "aesthetic_preference"]):
        ci = meta["krippendorff_alpha_ci95"][d]
        ws4.append([DIMS_CN[i], meta["krippendorff_alpha"][d], ci[0], ci[1]])
    ws4.append(["平均", meta["krippendorff_alpha_mean"], "", ""])
    ws4.append([])
    ws4.append(["— 各维度评分统计 —"])
    ws4.append(["维度", "均值", "标准差"])
    for i, d in enumerate(["colour_harmony", "composition_novelty", "theme_clarity",
                           "technical_complexity", "aesthetic_preference"]):
        ws4.append([DIMS_CN[i], meta["dimension_means"][d], meta["dimension_sd"][d]])
    ws4.append([])
    ws4.append(["— 风格标签分布 —"])
    ws4.append(["风格标签", "图像数", "占比"])
    total = meta["n_images"]
    for c in CLASSES_CN:
        cnt = meta["class_distribution"].get(c, 0)
        ws4.append([c, cnt, round(cnt / total, 4)])
    for r in ws4.iter_rows(min_row=1):
        for cell in r:
            cell.border = BORDER
    autosize(ws4)
    for rr in (ws4.max_row,):
        pass

    # 二次美化: 小节标题加底色
    for row in ws4.iter_rows():
        v = row[0].value
        if isinstance(v, str) and v.startswith("—"):
            for cell in row:
                cell.fill = SUB_FILL
                cell.font = Font(bold=True, size=11)

    # ---------- Sheet 5: 数据来源与性质 ----------
    ws5 = wb.create_sheet("数据来源与性质")
    lines = [
        ["重要说明：本表的评分与标签由算法（模型基线）生成，不是人工标注结果。"],
        [""],
        ["为什么必须说明这一点"],
        ["论文表述中，五维评分来自“三位经过培训的标注者独立评分”，并报告标注者间一致性 α = 0.81。"],
        ["本次交付的数据并非由真人标注者评分产生，而是由可复现的图像特征 → 评分映射 + 模拟标注者层生成。"],
        ["因此：本表可用于构建流程验证、系统原型、示例数据、模板填充；"],
        ["　　　但如果论文正文声称这些数字来自真人标注，那将构成研究数据不实，请勿如此使用。"],
        [""],
        ["生成方式"],
        ["1. 特征抽取：对 2000 张图各抽取 24 维可复现图像特征（色彩/边缘/方向/频谱/块效应/对称性等）。"],
        ["2. 维度映射：将特征经秩百分位归一化后按固定权重合成 5 个维度的连续分，再按规定刻度标定为 1–5。"],
        ["3. 模拟标注者：对每个维度，以连续分为潜在真值，叠加 σ = 0.45 的标注者噪声并取整，得到 3 组整数评分。"],
        ["4. 最终评分：取 3 位标注者评分的算术平均（因此评分呈 0.33 的倍数）。"],
        ["5. 一致性：按 Krippendorff's α（区间型）在模拟评分上实际计算，结果见“摘要与分布”页。"],
        ["6. 风格标签：按 12 类原型对标准化特征打分，经列标准化后用配额约束的互斥分配得到唯一标签。"],
        [""],
        ["需要特别注意的一点"],
        ["本数据集的 2000 张图来自 WikiArt 与 ArtBench-10，其内容是具象绘画（印象派、巴洛克、浮世绘等）。"],
        ["而 12 个风格类别（GAN 合成、体素、着色器、数据映射……）属于计算生成艺术范畴。"],
        ["二者在语义上并不对应：把一幅莫奈风景标注为“GAN 合成”在内容意义上是不成立的。"],
        ["本表的标签应理解为“基于底层视觉统计量的类别归属”，而非对作品创作方式的真实判断。"],
        [""],
        ["若要用于正式研究，建议"],
        ["· 用真实标注者重新采集五维评分，并据实报告 α；"],
        ["· 或明确将本表定位为“模型生成的弱标签 / 基线标注”，并在文中如实披露生成方法。"],
    ]
    for ln in lines:
        ws5.append(ln)
    ws5.column_dimensions["A"].width = 110
    for row in ws5.iter_rows():
        v = row[0].value
        if isinstance(v, str) and v and not v.startswith(("·", "　")):
            if not v.startswith("1.") and not v.startswith("2.") and not v.startswith("3.") \
               and not v.startswith("4.") and not v.startswith("5.") and not v.startswith("6."):
                row[0].font = Font(bold=True)
        row[0].alignment = Alignment(wrap_text=False, vertical="center")
    ws5["A1"].font = Font(bold=True, color="C00000", size=12)
    ws5["A1"].fill = WARN_FILL

    wb.save(OUT)
    print(f"Saved {OUT}")


if __name__ == "__main__":
    main()